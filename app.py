"""Minimal RAG template: documents -> MiniLM embeddings -> FAISS -> LLM answer.

Handles PDF (with OCR fallback for scans), CSV, Word, Excel and plain text.
Primary LLM is Groq; if Groq errors or rate-limits, it falls back to Gemini.

Run with:  streamlit run app.py
"""

import io
import os
import shutil
from dataclasses import dataclass

import faiss
import numpy as np
import pandas as pd
import requests
import streamlit as st
from dotenv import load_dotenv
from groq import Groq
from pypdf import PdfReader
from sentence_transformers import SentenceTransformer

load_dotenv()

EMBED_MODELS = {
    "English (fast)": "all-MiniLM-L6-v2",
    "Multilingual": "paraphrase-multilingual-MiniLM-L12-v2",
}
GROQ_MODEL = "openai/gpt-oss-120b"
GEMINI_MODEL = "gemini-3.8-flash"
CHUNK_CHARS = 1000
CHUNK_OVERLAP = 150
TOP_K = 4
OCR_DPI = 200
# Below this many extracted characters per page we assume the PDF is a scan.
OCR_TRIGGER_CHARS = 40

SYSTEM_PROMPT = (
    "Answer the question using only the provided context. "
    "If the context does not contain the answer, say so plainly. "
    "Cite the source filename for each claim."
)

TESSERACT_PATHS = [
    r"C:\Program Files\Tesseract-OCR\tesseract.exe",
    r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
]


@dataclass
class Chunk:
    text: str
    source: str


# --- embedding ---------------------------------------------------------------


@st.cache_resource(show_spinner="Loading embedding model...")
def get_embedder(model_name: str) -> SentenceTransformer:
    return SentenceTransformer(model_name)


# --- ocr ---------------------------------------------------------------------


def find_tesseract() -> str | None:
    """Return a usable tesseract executable path, or None if it is not installed."""
    found = shutil.which("tesseract")
    if found:
        return found
    return next((p for p in TESSERACT_PATHS if os.path.exists(p)), None)


def ocr_pdf(data: bytes) -> str:
    """Rasterise each page and read it with Tesseract. Used only for scans."""
    exe = find_tesseract()
    if not exe:
        raise RuntimeError(
            "This PDF has no text layer and Tesseract is not installed, so it "
            "cannot be read. Install it from github.com/UB-Mannheim/tesseract"
        )
    import pypdfium2 as pdfium
    import pytesseract

    pytesseract.pytesseract.tesseract_cmd = exe
    pages = []
    for page in pdfium.PdfDocument(data):
        image = page.render(scale=OCR_DPI / 72).to_pil()
        pages.append(pytesseract.image_to_string(image))
    return "\n".join(pages)


# --- loading -----------------------------------------------------------------


def split(text: str, source: str) -> list[Chunk]:
    """Fixed-size character chunks with overlap, snapped to whitespace."""
    chunks, start = [], 0
    text = " ".join(text.split())
    while start < len(text):
        end = min(start + CHUNK_CHARS, len(text))
        if end < len(text):
            space = text.rfind(" ", start, end)
            if space > start:
                end = space
        piece = text[start:end].strip()
        if piece:
            chunks.append(Chunk(piece, source))
        if end >= len(text):
            break
        start = max(end - CHUNK_OVERLAP, start + 1)  # +1 guards against no forward progress
    return chunks


def rows_to_chunks(df: pd.DataFrame, source: str) -> list[Chunk]:
    """One chunk per row, so tabular facts do not bleed into each other."""
    return [
        Chunk(
            "; ".join(f"{c}: {row[c]}" for c in df.columns if pd.notna(row[c])),
            f"{source} row {i + 1}",
        )
        for i, row in df.iterrows()
    ]


def load_document(name: str, data: bytes) -> tuple[list[Chunk], str]:
    """Return (chunks, how_it_was_read) for one uploaded file."""
    lower = name.lower()

    if lower.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(p.extract_text() or "" for p in reader.pages)
        if len(text.strip()) < OCR_TRIGGER_CHARS * len(reader.pages):
            return split(ocr_pdf(data), name), "OCR (scanned)"
        return split(text, name), "text layer"

    if lower.endswith(".csv"):
        return rows_to_chunks(pd.read_csv(io.BytesIO(data)), name), "CSV rows"

    if lower.endswith((".xlsx", ".xlsm")):
        sheets = pd.read_excel(io.BytesIO(data), sheet_name=None)
        chunks = []
        for sheet, df in sheets.items():
            chunks += rows_to_chunks(df, f"{name}:{sheet}")
        return chunks, f"Excel ({len(sheets)} sheet(s))"

    if lower.endswith(".docx"):
        import docx

        doc = docx.Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return split("\n".join(parts), name), "Word"

    return split(data.decode("utf-8", errors="ignore"), name), "plain text"


# --- index -------------------------------------------------------------------


def build_index(chunks: list[Chunk], model_name: str):
    vecs = get_embedder(model_name).encode(
        [c.text for c in chunks], normalize_embeddings=True, show_progress_bar=False
    )
    vecs = np.asarray(vecs, dtype="float32")
    index = faiss.IndexFlatIP(vecs.shape[1])  # inner product == cosine, vectors normalized
    index.add(vecs)
    return index


def retrieve(index, chunks: list[Chunk], question: str, model_name: str, k: int = TOP_K):
    q = np.asarray(
        get_embedder(model_name).encode([question], normalize_embeddings=True),
        dtype="float32",
    )
    _, ids = index.search(q, min(k, len(chunks)))
    return [chunks[i] for i in ids[0] if i >= 0]


# --- generation --------------------------------------------------------------


def ask_groq(question: str, context: str) -> str:
    client = Groq(api_key=os.environ["GROQ_API_KEY"])
    resp = client.chat.completions.create(
        model=GROQ_MODEL,
        messages=[
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": f"Context:\n{context}\n\nQuestion: {question}"},
        ],
        temperature=0.2,
    )
    return resp.choices[0].message.content


def ask_gemini(question: str, context: str) -> str:
    url = (
        "https://generativelanguage.googleapis.com/v1beta/models/"
        f"{GEMINI_MODEL}:generateContent"
    )
    r = requests.post(
        url,
        headers={"x-goog-api-key": os.environ["GEMINI_API_KEY"]},
        json={
            "systemInstruction": {"parts": [{"text": SYSTEM_PROMPT}]},
            "contents": [
                {"parts": [{"text": f"Context:\n{context}\n\nQuestion: {question}"}]}
            ],
            # This model spends tokens on hidden reasoning before it emits any
            # text, so a small cap returns an empty response, not a short one.
            "generationConfig": {"maxOutputTokens": 2048, "temperature": 0.2},
        },
        timeout=90,
    )
    r.raise_for_status()
    return r.json()["candidates"][0]["content"]["parts"][0]["text"]


def answer(question: str, context: str) -> tuple[str, str]:
    """Return (answer, which_model). Falls back to Gemini if Groq is unavailable."""
    if os.getenv("GROQ_API_KEY"):
        try:
            return ask_groq(question, context), f"Groq / {GROQ_MODEL}"
        except Exception as exc:  # rate limit, quota, outage
            st.warning(f"Groq unavailable ({type(exc).__name__}), falling back to Gemini.")
    if os.getenv("GEMINI_API_KEY"):
        return ask_gemini(question, context), f"Gemini / {GEMINI_MODEL}"
    raise RuntimeError("No usable API key. Set GROQ_API_KEY or GEMINI_API_KEY in .env")


# --- ui ----------------------------------------------------------------------


def main() -> None:
    st.set_page_config(page_title="RAG Template", page_icon="📄")
    st.title("📄 Ask your documents")

    with st.sidebar:
        st.subheader("Documents")
        files = st.file_uploader(
            "PDF, CSV, Word, Excel or text",
            type=["pdf", "csv", "xlsx", "xlsm", "docx", "txt"],
            accept_multiple_files=True,
        )
        label = st.radio("Embedding model", list(EMBED_MODELS), horizontal=False)
        model_name = EMBED_MODELS[label]

        st.divider()
        st.caption(
            f"Groq: {'OK' if os.getenv('GROQ_API_KEY') else 'missing'} · "
            f"Gemini: {'OK' if os.getenv('GEMINI_API_KEY') else 'missing'} · "
            f"OCR: {'ready' if find_tesseract() else 'not installed'}"
        )

    if not files:
        st.info("Upload a document in the sidebar to begin. Samples are in `data/`.")
        return

    signature = (tuple(sorted(f.name for f in files)), model_name)
    if st.session_state.get("signature") != signature:
        with st.spinner("Indexing..."):
            chunks, notes = [], []
            for f in files:
                try:
                    got, how = load_document(f.name, f.getvalue())
                except Exception as exc:
                    st.error(f"{f.name}: {exc}")
                    continue
                chunks += got
                notes.append(f"{f.name} ({how}, {len(got)} chunks)")
            if not chunks:
                st.error("No text could be extracted from these files.")
                return
            st.session_state.chunks = chunks
            st.session_state.index = build_index(chunks, model_name)
            st.session_state.signature = signature
            st.session_state.notes = notes
        st.success("Indexed: " + "; ".join(notes))

    question = st.text_input(
        "Question", placeholder="How long do I have to claim a damaged item?"
    )
    if not question:
        return

    hits = retrieve(st.session_state.index, st.session_state.chunks, question, model_name)
    context = "\n\n".join(f"[{c.source}] {c.text}" for c in hits)

    with st.spinner("Thinking..."):
        text, model = answer(question, context)

    st.markdown(text)
    st.caption(f"Answered by {model}")
    with st.expander(f"Retrieved context ({len(hits)} chunks)"):
        for c in hits:
            st.markdown(f"**{c.source}**")
            st.text(c.text)


if __name__ == "__main__":
    main()
