"""Turn files into Documents.

Each loader records how the text was obtained so the UI can show it and the
evaluator can tell an OCR miss apart from a retrieval miss.
"""

from __future__ import annotations

import io
import os
import shutil
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd

from .config import settings

TESSERACT_PATHS = [
    r"C:\Program Files\Tesseract-OCR\tesseract.exe",
    r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    "/usr/bin/tesseract",
    "/opt/homebrew/bin/tesseract",
]

TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".rst", ".log"}


@dataclass
class Document:
    """One logical unit of source text, before chunking."""

    text: str
    source: str
    how: str  # "text layer", "OCR (scanned)", "CSV rows", ...
    meta: dict = field(default_factory=dict)
    atomic: bool = False  # True => never split (a table row, say)


def find_tesseract() -> str | None:
    """Return a usable tesseract executable, or None if it is not installed."""
    return shutil.which("tesseract") or next(
        (p for p in TESSERACT_PATHS if os.path.exists(p)), None
    )


# --- pdf ---------------------------------------------------------------------


def _ocr_pdf(data: bytes) -> list[str]:
    exe = find_tesseract()
    if not exe:
        raise RuntimeError(
            "This PDF has no text layer and Tesseract is not installed, so it "
            "cannot be read. Install it: winget install UB-Mannheim.TesseractOCR"
        )
    import pypdfium2 as pdfium
    import pytesseract

    pytesseract.pytesseract.tesseract_cmd = exe
    pages = []
    for page in pdfium.PdfDocument(data):
        image = page.render(scale=settings.ocr_dpi / 72).to_pil()
        pages.append(pytesseract.image_to_string(image, lang=settings.ocr_lang))
    return pages


def load_pdf(name: str, data: bytes) -> list[Document]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [p.extract_text() or "" for p in reader.pages]
    how = "text layer"

    if sum(len(p.strip()) for p in pages) < settings.ocr_trigger_chars * max(len(pages), 1):
        pages = _ocr_pdf(data)
        how = "OCR (scanned)"

    # One Document per page keeps page numbers attached to every citation.
    return [
        Document(text=text, source=name, how=how, meta={"page": i + 1})
        for i, text in enumerate(pages)
        if text.strip()
    ]


# --- tables ------------------------------------------------------------------


def _frame_to_documents(df: pd.DataFrame, source: str, sheet: str | None = None) -> list[Document]:
    """One Document per row, plus a header summary.

    Rows are marked atomic so chunking never splits a record in half - this is
    what keeps tabular answers from mixing two customers together.
    """
    df = df.fillna("")
    label = f"{source}:{sheet}" if sheet else source
    columns = ", ".join(str(c) for c in df.columns)

    docs = [
        Document(
            text=f"Table {label} has {len(df)} rows with columns: {columns}.",
            source=label,
            how="table summary",
            meta={"kind": "schema"},
            atomic=True,
        )
    ]
    for i, row in df.iterrows():
        pairs = "; ".join(f"{c}: {row[c]}" for c in df.columns if str(row[c]).strip())
        if pairs:
            docs.append(
                Document(
                    text=pairs,
                    source=label,
                    how="table row",
                    meta={"row": int(i) + 1},
                    atomic=True,
                )
            )
    return docs


def load_csv(name: str, data: bytes) -> list[Document]:
    return _frame_to_documents(pd.read_csv(io.BytesIO(data)), name)


def load_excel(name: str, data: bytes) -> list[Document]:
    sheets = pd.read_excel(io.BytesIO(data), sheet_name=None)
    docs: list[Document] = []
    for sheet, df in sheets.items():
        docs += _frame_to_documents(df, name, sheet)
    return docs


# --- word / text -------------------------------------------------------------


def load_docx(name: str, data: bytes) -> list[Document]:
    import docx

    doc = docx.Document(io.BytesIO(data))
    docs = [
        Document(text=p.text, source=name, how="Word")
        for p in doc.paragraphs
        if p.text.strip()
    ]
    # Word tables become one atomic Document per row, same as a CSV.
    for t_i, table in enumerate(doc.tables, 1):
        rows = [[c.text.strip() for c in r.cells] for r in table.rows]
        if not rows:
            continue
        header, body = rows[0], rows[1:]
        for r_i, row in enumerate(body, 1):
            pairs = "; ".join(f"{h}: {v}" for h, v in zip(header, row) if v)
            if pairs:
                docs.append(
                    Document(
                        text=pairs,
                        source=name,
                        how="Word table",
                        meta={"table": t_i, "row": r_i},
                        atomic=True,
                    )
                )
    return docs


def load_text(name: str, data: bytes) -> list[Document]:
    return [Document(text=data.decode("utf-8", errors="ignore"), source=name, how="plain text")]


# --- dispatch ----------------------------------------------------------------


def load_bytes(name: str, data: bytes) -> list[Document]:
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        return load_pdf(name, data)
    if suffix == ".csv":
        return load_csv(name, data)
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        return load_excel(name, data)
    if suffix == ".docx":
        return load_docx(name, data)
    if suffix in TEXT_SUFFIXES:
        return load_text(name, data)
    raise ValueError(f"Unsupported file type: {suffix or name}")


def load_path(path: str | Path) -> list[Document]:
    path = Path(path)
    return load_bytes(path.name, path.read_bytes())


def load_dir(directory: str | Path) -> list[Document]:
    """Load every supported file in a directory, skipping ones that fail."""
    docs: list[Document] = []
    for path in sorted(Path(directory).iterdir()):
        if not path.is_file():
            continue
        try:
            docs += load_path(path)
        except (ValueError, RuntimeError):
            continue
    return docs
